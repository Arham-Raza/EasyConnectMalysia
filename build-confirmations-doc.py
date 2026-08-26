from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "EasyConnect-Pending-Confirmations.docx"
BLUE = RGBColor(41, 171, 226)
NAVY = RGBColor(0, 27, 58)
MUTED = RGBColor(71, 85, 105)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = NAVY
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
header.text = "EASY CONNECT SOLUTIONS  |  WEBSITE CONTENT REVIEW"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
for run in header.runs:
    run.font.name = "Calibri"; run.font.size = Pt(8); run.font.color.rgb = MUTED

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = footer.add_run("Internal working document  |  ")
run.font.size = Pt(8); run.font.color.rgb = MUTED
field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(4)
r = p.add_run("PENDING CONFIRMATIONS")
r.bold = True; r.font.name = "Calibri"; r.font.size = Pt(24); r.font.color.rgb = NAVY
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Easy Connect Solutions Malaysia website content and launch checklist")
r.font.size = Pt(12); r.font.color.rgb = MUTED

callout = doc.add_paragraph()
callout.paragraph_format.left_indent = Inches(.2)
callout.paragraph_format.right_indent = Inches(.2)
callout.paragraph_format.space_after = Pt(14)
shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "EAF7FC"); callout._p.get_or_add_pPr().append(shd)
r = callout.add_run("Purpose: "); r.bold = True; r.font.color.rgb = NAVY
callout.add_run("Confirm only the business facts that cannot be inferred safely. The website can continue development while these items remain unpublished or neutrally worded.")

sections = [
    ("1. Retail and online-store information", [
        "Full retail address for the Contact page, including postcode and unit/shop number.",
        "Google Maps share URL or embed location for the retail address.",
        "Official TechRevelo online-store URL, including whether the primary link should go to Shopee or another storefront.",
        "Short TechRevelo brand description approved for the Products and Contact pages.",
    ]),
    ("2. Customer proof and case studies", [
        "Approved customer testimonials with customer name, role, organization and permission to publish.",
        "Google review profile/link and the exact reviews approved for reuse.",
        "At least two completed projects suitable for case studies: client type, problem, scope, quantity, delivery and measurable outcome.",
        "Logos that Easy Connect has permission to display as customers or partners.",
    ]),
    ("3. Compliance and service evidence", [
        "Confirm whether data wiping is performed directly by Easy Connect or by a named qualified partner.",
        "Confirm whether NIST SP 800-88 and/or DoD 5220.22-M can be claimed, and provide the supporting process or partner evidence.",
        "Confirm whether Certificates of Data Destruction and chain-of-custody reports are issued, and at what scope: per drive, per device or per project.",
        "Confirm the approved wording for PDPA alignment/compliance.",
        "Confirm the downstream e-waste recovery partner and whether DOE-licensed claims can be published.",
    ]),
    ("4. Commercial terms", [
        "Warranty term by product category and grade, including whether 6-12 months is standard or optional.",
        "Corporate replacement/swap service: availability, geography, exclusions and the approved response time.",
        "Approved quotation turnaround and delivery-time wording.",
        "Minimum order quantities for bulk procurement and institutional pricing.",
        "Approved savings claims, including the evidence for 30%, 40% or other comparisons against new equipment.",
    ]),
    ("5. Product and inventory details", [
        "Final supported Apple range for the Editing page: MacBook Pro generations, iMac models and typical configurations.",
        "Confirm which gaming GPU generations are currently supportable in copy: RTX 30, 40 or newer series.",
        "Confirm whether Toshiba, Acer, Chromebooks and specific workstation models should remain in the permanent product copy.",
        "Approved grading definitions for Grade A and any other grades offered.",
    ]),
    ("6. Marketing and operations", [
        "Decision on WhatsApp chatbot implementation. The current site uses a direct WhatsApp link to +6011 2090 6561.",
        "Preferred CMS and publishing workflow for future blog articles; the current release keeps the static Resources structure.",
        "LinkedIn company-page URL when available; the current link is generic LinkedIn.",
        "Final scope and timing for Malay and Chinese translations across long-form inner-page content.",
    ]),
]

for heading, items in sections:
    doc.add_heading(heading, level=1)
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(.5)
        p.paragraph_format.first_line_indent = Inches(-.25)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(item)

doc.add_heading("Current implementation decisions", level=1)
for item in [
    "Remarketing has been removed from current navigation and search indexing because the audit explicitly requested its removal.",
    "Unverified guarantees and certifications are being replaced with neutral process wording until evidence is supplied.",
    "The contact form submits to info@easyconnect.my and the WhatsApp number remains +6011 2090 6561.",
    "No fabricated testimonials, client names, maps, retail addresses or store URLs will be published.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(.5)
    p.paragraph_format.first_line_indent = Inches(-.25)
    p.paragraph_format.space_after = Pt(8)
    p.add_run(item)

doc.save(OUT)
print(OUT)
