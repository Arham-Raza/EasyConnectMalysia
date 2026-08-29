from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "EasyConnect-Team-Information-Required.docx"
NAVY = RGBColor(0, 27, 58)
BLUE = RGBColor(41, 171, 226)
MUTED = RGBColor(71, 85, 105)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.85)
section.left_margin = section.right_margin = Inches(0.9)
section.header_distance = section.footer_distance = Inches(0.45)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = NAVY
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 14, 7),
    ("Heading 2", 12.5, NAVY, 10, 5),
]:
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = header.add_run("EASY CONNECT SOLUTIONS  |  WEBSITE INFORMATION REQUEST")
r.font.name = "Calibri"; r.font.size = Pt(8); r.font.color.rgb = MUTED

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("Team review  |  Page ")
r.font.size = Pt(8); r.font.color.rgb = MUTED
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

title = doc.add_paragraph()
title.paragraph_format.space_after = Pt(3)
r = title.add_run("INFORMATION REQUIRED")
r.bold = True; r.font.name = "Calibri"; r.font.size = Pt(24); r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(14)
r = sub.add_run("Easy Connect Solutions Malaysia website - team response document")
r.font.size = Pt(12); r.font.color.rgb = MUTED

callout = doc.add_paragraph()
callout.paragraph_format.left_indent = Inches(.18)
callout.paragraph_format.right_indent = Inches(.18)
callout.paragraph_format.space_after = Pt(12)
shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "EAF7FC")
callout._p.get_or_add_pPr().append(shd)
r = callout.add_run("Purpose: "); r.bold = True
callout.add_run("Please provide or approve the items below so temporary wording can be replaced with final verified information. Items marked Priority 1 affect launch credibility or legal/compliance claims.")

def add_item(number, prompt, guidance=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.36)
    p.paragraph_format.first_line_indent = Inches(-.36)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{number}. ")
    r.bold = True; r.font.color.rgb = BLUE
    r = p.add_run(prompt)
    r.bold = True
    if guidance:
        g = doc.add_paragraph(guidance)
        g.paragraph_format.left_indent = Inches(.36)
        g.paragraph_format.space_after = Pt(2)
        for run in g.runs: run.font.color.rgb = MUTED; run.font.size = Pt(9.5)
    response = doc.add_paragraph("Response / link / attachment: _________________________________________________")
    response.paragraph_format.left_indent = Inches(.36)
    response.paragraph_format.space_after = Pt(8)
    for run in response.runs: run.font.color.rgb = MUTED; run.font.size = Pt(9)

doc.add_heading("Priority 1 - Required before final launch", level=1)

doc.add_heading("A. Company, retail and online-store details", level=2)
add_item(1, "Full legal company name and registration number", "Needed for final legal-policy and commercial wording.")
add_item(2, "Complete retail address", "Include unit/shop number, street, postcode, city and state.")
add_item(3, "Google Maps location", "Provide the Google Maps share URL or approved map embed location.")
add_item(4, "Official TechRevelo store URL", "Confirm whether the primary link should open Shopee, another marketplace or a standalone store.")
add_item(5, "Approved TechRevelo description", "One or two sentences explaining its relationship with Easy Connect Solutions.")

doc.add_heading("B. ITAD, data security and e-waste evidence", level=2)
add_item(6, "Who performs data wiping?", "Confirm whether it is completed by Easy Connect directly or by a named service partner.")
add_item(7, "Approved data-sanitization standards", "Confirm whether NIST SP 800-88, DoD 5220.22-M or another standard can be stated publicly. Attach supporting evidence.")
add_item(8, "Data-destruction documentation", "Confirm whether certificates and chain-of-custody reports are issued per drive, per device or per project.")
add_item(9, "Approved PDPA wording", "Confirm whether the website may say compliant, aligned, designed to support compliance, or another approved phrase.")
add_item(10, "E-waste recovery partner and licensing", "Provide partner name, DOE/licensing evidence and approved public wording.")

doc.add_heading("C. Warranty and commercial commitments", level=2)
add_item(11, "Warranty terms by product category", "Confirm standard duration, optional extensions, exclusions, batteries and claim procedure.")
add_item(12, "Replacement or swap service", "Confirm whether a corporate swap service exists, where it applies and the approved response time.")
add_item(13, "Quotation and delivery turnaround", "Provide realistic public wording and any exclusions based on stock or location.")
add_item(14, "Minimum order quantities", "Confirm thresholds for bulk procurement, education and corporate pricing.")
add_item(15, "Approved savings claims", "Provide evidence and approved comparison wording for any 30%, 40% or other savings figure.")

doc.add_heading("Priority 2 - Required to replace temporary content", level=1)

doc.add_heading("D. Reviews, customers and case studies", level=2)
add_item(16, "Approved customer testimonials", "For each: quote, customer name, role, organization and written permission to publish.")
add_item(17, "Google review link and approved reviews", "Provide the business profile URL and identify the exact reviews that may be reused.")
add_item(18, "Two or more completed case studies", "For each: customer type, problem, project scope, quantity, delivery approach and measurable outcome.")
add_item(19, "Customer or partner logos", "Provide high-resolution files and confirmation that Easy Connect has permission to display them.")

doc.add_heading("E. Product and inventory confirmation", level=2)
add_item(20, "Approved grading definitions", "Confirm Grade A and any other grades, including cosmetic, display and battery expectations.")
add_item(21, "Apple editing range", "Confirm supported MacBook Pro generations, iMac models and typical RAM/storage configurations.")
add_item(22, "Gaming range", "Confirm supported brands, current GPU generations and whether RTX 30/40-series wording is appropriate.")
add_item(23, "Permanent product brands and models", "Confirm whether Acer, Toshiba, Chromebooks and named workstation models should remain in evergreen copy.")

doc.add_heading("Priority 3 - Marketing and future development", level=1)

doc.add_heading("F. Marketing operations", level=2)
add_item(24, "LinkedIn company-page URL", "The current website uses the generic LinkedIn homepage until a company URL is supplied.")
add_item(25, "WhatsApp chatbot decision", "Confirm direct WhatsApp only, or provide the preferred chatbot platform and required conversation flow.")
add_item(26, "Blog CMS decision", "Confirm preferred CMS and approval workflow when dynamic articles are introduced later.")
add_item(27, "Full translation scope", "Confirm whether long-form Product, Solution, Resource and Legal pages require full Malay and Chinese translations.")

doc.add_heading("Already confirmed", level=1)
for item in [
    "Website domain: easyconnect.my",
    "Contact email: info@easyconnect.my",
    "Phone and WhatsApp: +60 13-532 2733",
    "Base location: Shah Alam, Selangor, Malaysia",
    "Facebook and Instagram links supplied",
    "Remarketing page removed",
    "Blog CMS will be considered later",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(.5)
    p.paragraph_format.first_line_indent = Inches(-.25)
    p.paragraph_format.space_after = Pt(5)
    p.add_run(item)

doc.save(OUT)
print(OUT)
